"""Export the HTML motion deck to PowerPoint.

Renders each slide's final frame with headless Chrome (?static=1), then builds a 16:9 .pptx with a
fade/morph transition on every slide and the storyboard lines as speaker notes.

    python docs/build_deck.py [http://localhost:8666]
"""
from __future__ import annotations

import subprocess
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.util import Emu

HERE = Path(__file__).resolve().parent
OUT_IMG = HERE / "img" / "deck"
CHROME = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"
BASE = sys.argv[1] if len(sys.argv) > 1 else "http://localhost:8666"

NOTES = [
    "Every incident is different. Every dashboard is the same. This is WRAITH — the war room that builds itself.",
    "When something goes wrong you get a static dashboard built months ago for a different crisis, and a wall of unverified chatter.",
    "So we stopped building dashboards. WRAITH compiles one — for this incident — while the model is still thinking.",
    "A CrewAI Flow routes the incident. Three agents gather intel with partial knowledge. A Commander who only sees their briefs decides. Parasail streams a schema-locked interface over a WebSocket, and One acts on it.",
    "No agent sees everything — exactly like a real incident team. The Commander resolves their disagreement under doctrine.",
    "Every partner is load-bearing. Switch one off live and the station visibly breaks.",
    "Same engine, different incident, different interface. A judge can type any incident.",
    "The same swarm plays the Track 1 venue feed — evacuate, evacuate, and it refuses to escalate the human-crowd control tick. Zero false alarms.",
    "Halloween is the demo. Incidents are the market: SRE on-call, venues and events, emergency operations.",
    "We picked this idea by running 100 AI judges over 10 ideas — it won all 100 first-place votes. Open source, one command to run. WRAITH.",
]

P = "http://schemas.openxmlformats.org/presentationml/2006/main"
MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
P159 = "http://schemas.microsoft.com/office/powerpoint/2015/09/main"


def render() -> None:
    OUT_IMG.mkdir(parents=True, exist_ok=True)
    for n in range(1, 11):
        for attempt in range(3):  # headless Chrome occasionally exits early; a retry is enough
            res = subprocess.run([CHROME, "--headless=new", "--disable-gpu", "--hide-scrollbars", "--virtual-time-budget=5000",
                                  f"--screenshot={OUT_IMG / f'slide-{n}.png'}", "--window-size=1920,1080",
                                  f"{BASE}/docs/deck/index.html?static=1#{n}"], capture_output=True)
            if res.returncode == 0 and (OUT_IMG / f"slide-{n}.png").exists():
                break
        png = OUT_IMG / f"slide-{n}.png"
        subprocess.run(["sips", "-s", "format", "jpeg", "-s", "formatOptions", "82", str(png), "--out",
                        str(png.with_suffix(".jpg"))], check=True, capture_output=True)
        png.unlink()


def add_transition(slide, kind: str) -> None:
    """Morph where PowerPoint supports it, fade everywhere else."""
    sld = slide._element
    for old in sld.findall(f"{{{P}}}transition"):
        sld.remove(old)
    xml = (f'<mc:AlternateContent xmlns:mc="{MC}" xmlns:p="{P}" xmlns:p159="{P159}">'
           f'<mc:Choice Requires="p159"><p:transition spd="slow" p14:dur="1200" xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main">'
           f'{"<p159:morph option=\"byObject\"/>" if kind == "morph" else "<p:fade/>"}</p:transition></mc:Choice>'
           f'<mc:Fallback><p:transition spd="slow"><p:fade/></p:transition></mc:Fallback></mc:AlternateContent>')
    node = etree.fromstring(xml)
    clr = sld.find(f"{{{P}}}clrMapOvr")
    if clr is not None:
        clr.addnext(node)
    else:
        sld.find(f"{{{P}}}cSld").addnext(node)


def build() -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    blank = prs.slide_layouts[6]
    for n in range(1, 11):
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(OUT_IMG / f"slide-{n}.jpg"), 0, 0, width=prs.slide_width, height=prs.slide_height)
        slide.notes_slide.notes_text_frame.text = NOTES[n - 1]
        add_transition(slide, "morph" if n > 1 else "fade")
    out = HERE / "WRAITH-deck.pptx"
    prs.save(out)
    return out


if __name__ == "__main__":
    render()
    print(build())
